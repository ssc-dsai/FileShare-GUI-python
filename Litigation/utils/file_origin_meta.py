# Litigation/utils/file_origin_meta.py
"""Best-effort Created By / Last Modified By and filesystem timestamps."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path


def filesystem_times(path: Path) -> dict:
    st = path.stat()
    ctime = getattr(st, "st_ctime", st.st_mtime)
    return {
        "file_created": datetime.fromtimestamp(ctime).isoformat(timespec="seconds"),
        "file_modified": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
    }


def office_authors(path: Path) -> dict:
    out = {"created_by": "", "last_modified_by": ""}
    suffix = path.suffix.lower()

    if suffix == ".docx":
        try:
            from docx import Document
            props = Document(str(path)).core_properties
            out["created_by"] = (props.author or "")[:200]
            out["last_modified_by"] = (props.last_modified_by or "")[:200]
        except Exception:
            pass
        return out

    if suffix == ".pdf":
        try:
            import fitz
            doc = fitz.open(path)
            meta = doc.metadata or {}
            doc.close()
            out["created_by"] = (meta.get("author") or "")[:200]
            out["last_modified_by"] = (meta.get("producer") or "")[:200]
        except Exception:
            pass
        return out

    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            props = wb.properties
            out["created_by"] = (props.creator or "")[:200]
            out["last_modified_by"] = (props.lastModifiedBy or "")[:200]
            wb.close()
        except Exception:
            pass
        return out

    return out


def origin_metadata(path: Path) -> dict:
    data = {"created_by": "", "last_modified_by": "", "file_created": "", "file_modified": ""}
    try:
        if path.is_file():
            data.update(filesystem_times(path))
            data.update(office_authors(path))
    except Exception:
        pass
    return data